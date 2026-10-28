"""Optional: regenerate B's three editable presentation slides (python-pptx).

Not a production dependency. Run the B browser tests first for the real screenshot.
"""
from pathlib import Path
import json

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
FONT = "Microsoft YaHei"
INK, MUTED, GREEN, BLUE, CREAM = "203F38", "60756C", "2B7258", "287C91", "F3F7F2"


def color(value):
    return RGBColor.from_string(value)


def box(slide, x, y, width, height, fill="FFFFFF", border=None, rounded=False):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(width), Inches(height))
    shape.fill.solid(); shape.fill.fore_color.rgb = color(fill)
    if border: shape.line.color.rgb = color(border)
    else: shape.line.fill.background()
    return shape


def text(slide, value, x, y, width, height, size=22, ink=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame; frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(.01)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line; paragraph.alignment = align; paragraph.space_after = Pt(8)
        for run in paragraph.runs:
            run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color(ink)
    return shape


def line(slide, x1, y1, x2, y2):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = color(BLUE); shape.line.width = Pt(2)


def player(slide, name, x, y):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x-.15), Inches(y-.15), Inches(.3), Inches(.3))
    shape.fill.solid(); shape.fill.fore_color.rgb = color(BLUE); shape.line.fill.background()
    text(slide, name, x-.5, y+.25, 1, .4, 16, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    edges = json.loads((ROOT / "data/processed/collaboration_edges.json").read_text())
    screenshot = ROOT / "docs/screenshots/hypergraph-b-main.png"
    if not screenshot.exists():
        raise SystemExit("Run tests/hypergraph.browser.mjs to capture the real graph first.")

    def slide(title, index):
        page = prs.slides.add_slide(prs.slide_layouts[6]); page.background.fill.solid(); page.background.fill.fore_color.rgb = color(CREAM)
        text(page, "FOOTBALL HYPERGRAPH  /  MEMBER B", .6, .3, 11.5, .3, 12, GREEN, True)
        text(page, title, .6, .85, 12.1, .8, 29, bold=True)
        text(page, f"StatsBomb Open Data · match 3869685 | Codex 辅助实现、测试与材料整理 | B · {index}/3", .6, 7.0, 12.1, .25, 10, MUTED)
        return page

    first = slide("把一次多人配合，保留为一个集合", 1)
    box(first, .6, 1.95, 5.85, 4.65, rounded=True)
    box(first, 6.85, 1.95, 5.85, 4.65, rounded=True)
    text(first, "两两关系示意", 1.0, 2.25, 5, .5, 21, bold=True)
    text(first, "同一片段的完整集合", 7.25, 2.25, 5, .5, 21, bold=True)
    left = [(2.1, 3.5), (4.9, 3.5), (3.5, 5.2)]
    for a, b in [(0, 1), (1, 2), (2, 0)]: line(first, *left[a], *left[b])
    for name, point in zip("ABC", left): player(first, name, *point)
    envelope = first.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(7.9), Inches(3.0), Inches(3.9), Inches(2.85))
    envelope.fill.solid(); envelope.fill.fore_color.rgb = color("E0EFE6"); envelope.line.color.rgb = color(GREEN); envelope.line.width = Pt(2)
    for name, point in zip("ABC", [(8.5, 3.5), (11.1, 3.5), (9.8, 5.2)]): player(first, name, *point)
    text(first, "每条连线关联两人", 1.0, 5.95, 5, .4, 17, MUTED)
    text(first, "e = {A, B, C} · 本项目至少三人", 7.25, 5.95, 5, .4, 17, GREEN)
    first.notes_slide.notes_text_frame.text = "建议45秒。" + "我的部分解决的是，怎样把一段多人配合作为一个整体保留下来。普通传球图中的一条线只连接两个人；如果只展示两两关系，读者不容易知道哪些球员属于同一个事件片段。我们用一条包络表示完整成员集合，本项目要求至少三名不同球员。节点是球员，超边是有来源依据的片段，并且可以点击回到原始事件。这里是概念示意，关系布局不代表球场位置。"

    second = slide("先确定证据，再构建有界传球窗口", 2)
    text(second, str(len(edges)), .65, 1.85, 3.6, 1.0, 56, GREEN, True)
    text(second, "真实原始协作片段", .7, 2.95, 4.8, .5, 22)
    text(second, "最多 4 次传球 / 20 秒 / 5 人", .7, 3.6, 5.2, .55, 23, GREEN, True)
    text(second, "按比赛顺序 → 同队、同控球\n成功传球相连 → 三人以上才输出\n每条 weight=1，保留时间和源事件", .7, 4.5, 5.5, 1.75, 20)
    chart_data = CategoryChartData(); chart_data.categories = ["三人", "四人", "五人"]
    chart_data.add_series("原始片段数", [sum(len(edge["player_ids"]) == n for edge in edges) for n in [3, 4, 5]])
    chart = second.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.6), Inches(2.1), Inches(5.9), Inches(3.8), chart_data).chart
    chart.has_legend = False; chart.chart_style = 10
    chart.value_axis.minimum_scale = 0; chart.value_axis.maximum_scale = 300; chart.value_axis.major_unit = 100
    chart.category_axis.tick_labels.font.name = FONT; chart.category_axis.tick_labels.font.size = Pt(16)
    chart.value_axis.tick_labels.font.name = FONT; chart.value_axis.tick_labels.font.size = Pt(12)
    plot = chart.plots[0]; plot.has_data_labels = True; plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.name = FONT; plot.data_labels.font.size = Pt(16)
    chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb = color(GREEN)
    text(second, "窗口可以重叠；片段数不等于独立进攻次数。", 6.75, 6.05, 5.6, .65, 16, MUTED)
    second.notes_slide.notes_text_frame.text = "建议60秒。" + "A已经提供真实比赛的规范化事件。我按阶段和事件序号恢复顺序，只使用常规时间与加时赛。传球必须明确成功，下一位发起者要接上上一位传球对象；换控球、失败传球和时间异常都会切断链。每次最多四次传球、二十秒、五名球员，去重达到三人才生成超边。默认得到五百三十一条原始片段，三人、四人、五人的分布见图。每条保存时间和源事件，权重是一；重复组合在展示时才合并。窗口会重叠，因此片段数不是独立进攻次数。参数可配置，报告比较了两次、三次、四次传球的结果。"

    third = slide("让球星、事件证据与时间在同一个主图联动", 3)
    box(third, .6, 1.9, 7.8, 4.85, rounded=True)
    with Image.open(screenshot) as image: width, height = image.size
    scale = min(7.45 / width, 4.5 / height)
    image_width, image_height = width * scale, height * scale
    third.shapes.add_picture(str(screenshot), Inches(.78 + (7.45-image_width)/2), Inches(2.06 + (4.5-image_height)/2), width=Inches(image_width), height=Inches(image_height))
    text(third, "选择 Messi", 8.85, 2.15, 3.9, .5, 25, GREEN, True)
    text(third, "81 个协作片段\n点击组合，展开源事件", 8.85, 2.85, 3.9, 1.1, 20)
    text(third, "切换对抗 / 调整时间", 8.85, 4.2, 3.9, .65, 23, GREEN, True)
    text(third, "同一 SVG · 虚线表示推断\n拖拽 / 缩放 / 键盘可用\nTop 12/24 明示显示上限", 8.85, 5.0, 3.9, 1.5, 18)
    third.notes_slide.notes_text_frame.text = "建议65秒。现场依次选择Messi、点击组合并展开事件、切Competition、缩小时间范围。" + "选择Messi后，主图高亮相关组合，C面板显示八十一个协作片段。默认仅显示前十二个完整组合，所有原始片段仍保留在统计和证据列表。点击包络可以看到全部成员、时间和事件ID，也能从图中选择球员反向更新面板。主图支持拖拽、缩放、平移和键盘选择。对抗模式复用同一个SVG，使用D的数据和虚线样式，明确这是事件窗口推断。球队、时间和球员通过共享状态联动。我的交付包括协作算法、主图、测试及真实案例。AI辅助实现和验证，成员本人还需审阅和计时排练。数据来源：StatsBomb Open Data，固定提交4b73468fc5b0f1950f9f66fada70ad3a4f9327cb。D3 7.9.0，ISC许可证。"
    output = ROOT / "docs/slides/part_B.pptx"; output.parent.mkdir(parents=True, exist_ok=True); prs.save(output)
    loaded = Presentation(output)
    assert len(loaded.slides) == 3
    for page in loaded.slides:
        assert page.has_notes_slide
        for shape in page.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= loaded.slide_width + 100
            assert shape.top + shape.height <= loaded.slide_height + 100
    print(f"Generated {output}: 3 slides, editable text/diagrams/chart, geometry checked.")


if __name__ == "__main__":
    main()
